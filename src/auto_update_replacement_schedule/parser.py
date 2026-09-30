# parser.py
import io
import re
from dataclasses import dataclass
from datetime import date, time

import dateparser
from docx import Document
from loguru import logger


@dataclass(slots=True)
class ParsedLessonDTO:
    date: date
    group_name: str
    class_start: time
    class_end: time
    subject_name: str
    class_type_name: str
    audience_name: str


def parse_docx_bytes(
    file_bytes: bytes, target_date: date | None = None
) -> tuple[date | None, list[ParsedLessonDTO]]:
    """
    Синхронный парсинг DOCX. Выполняется в отдельном потоке (asyncio.to_thread).
    Возвращает найденную дату и список распарсенных сырых строк замен.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except ValueError:
        logger.warning("Не удалось загрузить файл, возможно он не является .docx")
        return None, []
    schedule_date = target_date

    # 1. Поиск даты в документе, если не передана
    if not schedule_date:
        for table in doc.tables:
            for row in table.rows:
                if not row.cells:
                    continue
                text = row.cells[0].text.strip()
                if "(" in text and ")" in text:
                    try:
                        date_part = text.split("(")[1].split(")")[0].strip()
                        dt = dateparser.parse(
                            " ".join(date_part.split()[:3]), languages=["ru"]
                        )
                        if dt:
                            schedule_date = dt.date()
                            break
                    except Exception:
                        continue
            if schedule_date:
                break

    if not schedule_date:
        return None, []

    parsed_lessons: list[ParsedLessonDTO] = []

    # 2. Парсинг строк
    for table in doc.tables:
        for row in table.rows:
            cells_text = [
                [line.strip() for line in cell.text.strip().split("\n") if line.strip()]
                for cell in row.cells
            ]

            if len(cells_text) < 4 or not cells_text[0]:
                continue

            raw_group_name = cells_text[0][0]

            try:
                # Предмет и тип пары
                subject_cell = cells_text[2][0] if cells_text[2] else ""
                if " (" in subject_cell:
                    subject_parts = subject_cell.split(" (")
                    subject_name = subject_parts[0].strip()
                    class_type_name = (
                        subject_parts[1].replace(")", "").replace(",", "").strip()
                    )
                else:
                    subject_name = subject_cell.strip()
                    class_type_name = ""

                # Время пары
                if len(cells_text[1]) < 2:
                    continue
                time_raw = re.sub(r"\s+", "", cells_text[1][1])
                class_time = re.split(r"[-–—]", time_raw)
                if len(class_time) < 2:
                    continue

                h_start, m_start = re.split(r"[:.]", class_time[0])
                class_start_at = time(int(h_start), int(m_start))

                h_end, m_end = re.split(r"[:.]", class_time[1])
                class_end_at = time(int(h_end), int(m_end))

                # Аудитория
                audience_name = cells_text[3][0] if cells_text[3] else ""

                parsed_lessons.append(
                    ParsedLessonDTO(
                        date=schedule_date,
                        group_name=raw_group_name,
                        class_start=class_start_at,
                        class_end=class_end_at,
                        subject_name=subject_name,
                        class_type_name=class_type_name,
                        audience_name=audience_name,
                    )
                )
            except Exception as e:
                logger.warning(f"Ошибка парсинга строки: {e}")
                continue

    return schedule_date, parsed_lessons
